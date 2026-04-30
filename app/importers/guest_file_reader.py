from __future__ import annotations

import csv
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


SUPPORTED_FILE_EXTENSIONS = {".csv", ".xlsx"}


class GuestFileReadError(Exception):
    """Raised when a guest file cannot be read."""


def read_guest_file(file_path: str | Path) -> tuple[list[dict[str, Any]], set[str]]:
    """
    Read guest data from a CSV or Excel file.

    Returns:
    - rows: list of dictionaries
    - columns: normalized column names found in the file
    """

    path = Path(file_path)

    if not path.exists():
        raise GuestFileReadError(f"File does not exist: {path}")

    if not path.is_file():
        raise GuestFileReadError(f"Path is not a file: {path}")

    extension = path.suffix.lower()

    if extension not in SUPPORTED_FILE_EXTENSIONS:
        raise GuestFileReadError(
            f"Unsupported file type '{extension}'. Supported types: CSV, XLSX"
        )

    if extension == ".csv":
        return _read_csv(path)

    if extension == ".xlsx":
        return _read_xlsx(path)

    raise GuestFileReadError(f"Unsupported file type: {extension}")


def _read_csv(path: Path) -> tuple[list[dict[str, Any]], set[str]]:
    try:
        with path.open("r", encoding="utf-8-sig", newline="") as file:
            reader = csv.DictReader(file)

            if reader.fieldnames is None:
                raise GuestFileReadError("CSV file has no header row")

            normalized_headers = [_normalize_column_name(name) for name in reader.fieldnames]
            columns = set(normalized_headers)

            rows: list[dict[str, Any]] = []

            for index, row in enumerate(reader, start=2):
                normalized_row = {
                    _normalize_column_name(key): value
                    for key, value in row.items()
                    if key is not None
                }

                normalized_row["_row_number"] = index
                rows.append(normalized_row)

            return rows, columns

    except UnicodeDecodeError as exc:
        raise GuestFileReadError(
            "Could not read CSV file. Please save it as UTF-8 CSV and try again."
        ) from exc
    except csv.Error as exc:
        raise GuestFileReadError("Could not parse CSV file") from exc


def _read_xlsx(path: Path) -> tuple[list[dict[str, Any]], set[str]]:
    try:
        workbook = load_workbook(
            filename=path,
            read_only=True,
            data_only=True,
        )
    except Exception as exc:
        raise GuestFileReadError("Could not open Excel file") from exc

    worksheet = workbook.active

    rows_iterator = worksheet.iter_rows(values_only=True)

    try:
        header_row = next(rows_iterator)
    except StopIteration as exc:
        raise GuestFileReadError("Excel file is empty") from exc

    if not header_row:
        raise GuestFileReadError("Excel file has no header row")

    normalized_headers = [_normalize_column_name(value) for value in header_row]
    columns = {column for column in normalized_headers if column}

    rows: list[dict[str, Any]] = []

    for excel_row_number, values in enumerate(rows_iterator, start=2):
        row_data: dict[str, Any] = {}

        for column_name, value in zip(normalized_headers, values):
            if not column_name:
                continue

            row_data[column_name] = value

        if _is_empty_row(row_data):
            continue

        row_data["_row_number"] = excel_row_number
        rows.append(row_data)

    return rows, columns


def _normalize_column_name(value: Any) -> str:
    if value is None:
        return ""

    return str(value).strip().lower()


def _is_empty_row(row: dict[str, Any]) -> bool:
    ignored_keys = {"_row_number"}

    for key, value in row.items():
        if key in ignored_keys:
            continue

        if value is not None and str(value).strip():
            return False

    return True